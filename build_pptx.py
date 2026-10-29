import os

PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
except ModuleNotFoundError:
    venv_candidates = [
        os.path.join(PROJECT_ROOT, "backend", "venv", "Scripts", "python.exe"),
        os.path.join(PROJECT_ROOT, "backend", "venv", "bin", "python"),
    ]
    venv_python = next((path for path in venv_candidates if os.path.exists(path)), None)

    if venv_python:
        print(
            "python-pptx is not available in the active interpreter. "
            f"Run the project virtual environment instead:\n{venv_python} {os.path.basename(__file__)}"
        )
    else:
        print(
            "python-pptx is required to generate the presentation. "
            "Install it with:\npython -m pip install python-pptx"
        )
    raise SystemExit(1)

def create_presentation():
    prs = Presentation()
    # Set 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_slide_layout = prs.slide_layouts[6]

    # Theme colors
    DARK_BG = RGBColor(9, 13, 22)         # #090d16
    CARD_BG = RGBColor(17, 24, 39)        # #111827
    CYAN_NEON = RGBColor(6, 182, 212)     # #06b6d4
    EMERALD = RGBColor(16, 185, 129)      # #10b981
    YELLOW = RGBColor(245, 158, 11)       # #f59e0b
    WHITE = RGBColor(255, 255, 255)
    LIGHT_GRAY = RGBColor(203, 213, 225)
    DARK_GRAY = RGBColor(148, 163, 184)

    slides_data = [
        {
            "type": "title",
            "title": "CYVANTIS",
            "subtitle": "Enterprise Gateway Intelligence & Security Nexus",
            "tagline": "Next-Gen AI-Driven Honeypot Threat Intelligence & Automated Anomaly Prevention Engine",
            "track": "IMPRENDITORE 5.0 | STARTUP SPRINT CHALLENGE 2026 (16-Hour Offline Hackathon)",
            "bullets": [
                "Primary Track: Safety & Security",
                "Secondary Track: AI & Emerging Technologies",
                "Special Track: Open Innovation"
            ]
        },
        {
            "type": "problem",
            "title": "01. THE PROBLEM STATEMENT",
            "subtitle": "Why Traditional Cyber Defense Fails Against Modern Intrusions",
            "cards": [
                {
                    "head": "💥 2.2M+ Daily Cyber Attacks",
                    "text": "Automated brute-force & botnet scripts attack enterprise cloud gateways continuously."
                },
                {
                    "head": "❌ Legacy Firewall Blindspots",
                    "text": "Static rule-based firewalls only match known signatures and fail against Zero-Day exploits."
                },
                {
                    "head": "⚠️ Zero Real-Time Intelligence",
                    "text": "SOCs lack live honeypot decoy telemetry to trap attackers, inspect payloads, and train predictive AI."
                }
            ]
        },
        {
            "type": "solution",
            "title": "02. THE SOLUTION — CYVANTIS NEXUS",
            "subtitle": "Transforming Passive Enterprise Defense into Proactive AI Threat Prevention",
            "cards": [
                {
                    "head": "🍯 Virtualized Honeypot Traps",
                    "text": "Deploys virtualized Cowrie SSH/Telnet decoy nodes to trap attackers cleanly without risking assets."
                },
                {
                    "head": "🤖 Dual AI ML Inference Engine",
                    "text": "Trained on 200,000 telemetry events to predict attack severity and flag zero-day anomalies in real time."
                },
                {
                    "head": "🖥️ Interactive SOC Terminal",
                    "text": "Provides analysts with a 60FPS Matrix-style CMD console to search passwords, audit URLs, and run AI inference."
                }
            ]
        },
        {
            "type": "features",
            "title": "03. KEY TECHNICAL INNOVATIONS",
            "subtitle": "High-Impact Core Capabilities Built During the 16-Hour Hackathon",
            "list": [
                "🤖 Supervised Threat Severity Classifier: Scikit-learn RandomForest predicting Low, Medium, High, Critical severity.",
                "🔍 Unsupervised Anomaly Outlier Detector: IsolationForest model pinpointing unknown zero-day attacks.",
                "🔑 Database Credential Search Engine ('creds'): Instant lookup of captured attacker usernames & passwords.",
                "🌐 Attacker Origin & Payload Recon ('top-ips', 'top-cmds'): Aggregates source IPs, GeoIP, ASNs & script payloads.",
                "🛡️ Defensive Website URL Security Audit ('audit-url'): Assesses SSL/TLS encryption, script droppers & header compliance."
            ]
        },
        {
            "type": "architecture",
            "title": "04. TECHNICAL ARCHITECTURE & STACK",
            "subtitle": "Production-Grade Full Stack Implementation",
            "columns": [
                {
                    "title": "⚛️ Frontend SOC Console",
                    "items": ["React 18 + TypeScript", "Vite 8 Build Engine", "TailwindCSS & Dark Styling", "60FPS HTML5 Matrix Canvas FX"]
                },
                {
                    "title": "🐍 Backend REST API Gateway",
                    "items": ["Python 3.11 + FastAPI (Async)", "Uvicorn ASGI Web Server", "SQLite 3 (cyvantis.db)", "SQLAlchemy 2.0 ORM"]
                },
                {
                    "title": "🤖 Machine Learning Pipeline",
                    "items": ["Scikit-Learn (v1.6+)", "Pandas & NumPy Vectors", "Joblib Model Serialization", "200,000 Telemetry Scale"]
                }
            ]
        },
        {
            "type": "metrics",
            "title": "05. MACHINE LEARNING MODEL PERFORMANCE",
            "subtitle": "Empirical Verification on 200,000 Honeypot Telemetry Events",
            "stats": [
                {"number": "200,000", "label": "Telemetry Dataset Scale"},
                {"number": "92.40%", "label": "Random Forest Test Accuracy"},
                {"number": "92.30%", "label": "Weighted F1 Score Balance"},
                {"number": "9.99%", "label": "Isolation Forest Anomaly Rate"}
            ]
        },
        {
            "type": "terminal",
            "title": "06. LIVE DEMONSTRATION & COMMAND SHOWCASE",
            "subtitle": "Interactive Cyber Command Sandbox (http://localhost:5173/ml-intelligence)",
            "commands": [
                {"cmd": "creds --query admin", "desc": "Searches captured attacker usernames & passwords from cyvantis.db"},
                {"cmd": "top-ips", "desc": "Ranks top attacking IP addresses, ASN network & country of origin"},
                {"cmd": "top-cmds", "desc": "Ranks top malicious shell script payloads executed in honeypots"},
                {"cmd": "audit-url <url>", "desc": "Defensive website URL security audit & header compliance check"},
                {"cmd": "pipeline", "desc": "Displays scikit-learn model architecture, parameters & joblib specs"}
            ]
        },
        {
            "type": "market",
            "title": "07. MARKET OPPORTUNITY & BUSINESS MODEL",
            "subtitle": "Commercial Viability & Go-To-Market Strategy",
            "cards": [
                {
                    "head": "🎯 Target Audience",
                    "text": "Enterprise SOCs, Cloud Infrastructure Providers, Managed Security Service Providers (MSSPs)."
                },
                {
                    "head": "💵 B2B SaaS Subscriptions",
                    "text": "Tiered monthly licensing per monitored enterprise gateway node."
                },
                {
                    "head": "📡 Threat Feed Licensing",
                    "text": "Selling real-time honeypot IOC & IP reputation API data feeds."
                }
            ]
        },
        {
            "type": "roadmap",
            "title": "08. FUTURE ROADMAP & SCALABILITY",
            "subtitle": "From Hackathon Prototype to Enterprise Production",
            "steps": [
                {"phase": "Phase 1 (16h Hackathon MVP)", "detail": "Full FastAPI backend, React SOC dashboard, 200k ML pipeline, interactive CMD terminal."},
                {"phase": "Phase 2 (Next 3 Months)", "detail": "Kubernetes Container Honeypot Cluster integration & automated eBPF firewall IP blocking."},
                {"phase": "Phase 3 (Next 6 Months)", "detail": "Multi-Cloud deployment (AWS / Azure / GCP) & LLM Threat Intelligence Summarizer."}
            ]
        },
        {
            "type": "conclusion",
            "title": "09. CONCLUSION & HACKATHON SUMMARY",
            "subtitle": "Transforming Passive Defense into Proactive AI Threat Prevention",
            "bullets": [
                "✓ Built fully working FastAPI Backend & React Frontend in 16-Hour Startup Sprint.",
                "✓ Trained Scikit-learn Random Forest (92.4% Acc) & Isolation Forest on 200k events.",
                "✓ Production-ready code, live Swagger docs (http://127.0.0.1:8000/api/v1/docs).",
                "✓ Ready for live Q&A & Demonstration!"
            ]
        }
    ]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # Add background rect
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()

        if data["type"] == "title":
            # Title Slide
            tx_box = slide.shapes.add_textbox(Inches(1), Inches(1.2), Inches(11.33), Inches(5))
            tf = tx_box.text_frame
            tf.word_wrap = True

            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.size = Pt(54)
            p.font.bold = True
            p.font.color.rgb = CYAN_NEON

            p2 = tf.add_paragraph()
            p2.text = data["subtitle"]
            p2.font.size = Pt(28)
            p2.font.bold = True
            p2.font.color.rgb = WHITE
            p2.space_before = Pt(10)

            p3 = tf.add_paragraph()
            p3.text = data["tagline"]
            p3.font.size = Pt(18)
            p3.font.color.rgb = LIGHT_GRAY
            p3.space_before = Pt(15)

            p4 = tf.add_paragraph()
            p4.text = data["track"]
            p4.font.size = Pt(16)
            p4.font.bold = True
            p4.font.color.rgb = EMERALD
            p4.space_before = Pt(30)

        elif data["type"] in ["problem", "solution", "market"]:
            # Title
            tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
            tf = tx_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = CYAN_NEON

            p2 = tf.add_paragraph()
            p2.text = data["subtitle"]
            p2.font.size = Pt(16)
            p2.font.color.rgb = LIGHT_GRAY

            # 3 Cards
            card_width = Inches(3.64)
            card_height = Inches(4.5)
            spacing = Inches(0.4)
            left_start = Inches(0.8)

            for i, card in enumerate(data["cards"]):
                card_left = left_start + i * (card_width + spacing)
                card_top = Inches(2.0)

                # Card box shape
                c_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_left, card_top, card_width, card_height)
                c_box.fill.solid()
                c_box.fill.fore_color.rgb = CARD_BG
                c_box.line.color.rgb = CYAN_NEON if data["type"] == "solution" else YELLOW

                ctf = c_box.text_frame
                ctf.word_wrap = True
                ctf.margin_left = Inches(0.3)
                ctf.margin_right = Inches(0.3)
                ctf.margin_top = Inches(0.4)

                cp = ctf.paragraphs[0]
                cp.text = card["head"]
                cp.font.size = Pt(20)
                cp.font.bold = True
                cp.font.color.rgb = WHITE

                cp2 = ctf.add_paragraph()
                cp2.text = card["text"]
                cp2.font.size = Pt(15)
                cp2.font.color.rgb = LIGHT_GRAY
                cp2.space_before = Pt(15)

        elif data["type"] == "features":
            # Title
            tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
            tf = tx_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = CYAN_NEON

            p2 = tf.add_paragraph()
            p2.text = data["subtitle"]
            p2.font.size = Pt(16)
            p2.font.color.rgb = LIGHT_GRAY

            # List
            tx_box2 = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.8))
            tf2 = tx_box2.text_frame
            tf2.word_wrap = True

            for idx, item in enumerate(data["list"]):
                p = tf2.paragraphs[0] if idx == 0 else tf2.add_paragraph()
                p.text = f"• {item}"
                p.font.size = Pt(18)
                p.font.color.rgb = WHITE
                p.space_before = Pt(14)

        elif data["type"] == "architecture":
            # Title
            tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
            tf = tx_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = CYAN_NEON

            # 3 Columns
            col_w = Inches(3.64)
            col_h = Inches(4.5)
            spacing = Inches(0.4)
            left_start = Inches(0.8)

            for i, col in enumerate(data["columns"]):
                c_left = left_start + i * (col_w + spacing)
                c_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, Inches(2.0), col_w, col_h)
                c_box.fill.solid()
                c_box.fill.fore_color.rgb = CARD_BG
                c_box.line.color.rgb = EMERALD

                ctf = c_box.text_frame
                ctf.word_wrap = True
                ctf.margin_left = Inches(0.3)
                ctf.margin_top = Inches(0.3)

                cp = ctf.paragraphs[0]
                cp.text = col["title"]
                cp.font.size = Pt(20)
                cp.font.bold = True
                cp.font.color.rgb = CYAN_NEON

                for it in col["items"]:
                    cp2 = ctf.add_paragraph()
                    cp2.text = f"✓ {it}"
                    cp2.font.size = Pt(16)
                    cp2.font.color.rgb = WHITE
                    cp2.space_before = Pt(12)

        elif data["type"] == "metrics":
            # Title
            tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
            tf = tx_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = CYAN_NEON

            # 4 Metric Boxes
            b_w = Inches(2.65)
            b_h = Inches(4.2)
            spacing = Inches(0.37)
            left_start = Inches(0.8)

            for i, st in enumerate(data["stats"]):
                s_left = left_start + i * (b_w + spacing)
                s_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, s_left, Inches(2.2), b_w, b_h)
                s_box.fill.solid()
                s_box.fill.fore_color.rgb = CARD_BG
                s_box.line.color.rgb = CYAN_NEON if i % 2 == 0 else EMERALD

                stf = s_box.text_frame
                stf.word_wrap = True
                stf.margin_top = Inches(0.8)

                sp = stf.paragraphs[0]
                sp.text = st["number"]
                sp.font.size = Pt(36)
                sp.font.bold = True
                sp.font.color.rgb = EMERALD if i % 2 == 0 else CYAN_NEON
                sp.alignment = PP_ALIGN.CENTER

                sp2 = stf.add_paragraph()
                sp2.text = st["label"]
                sp2.font.size = Pt(16)
                sp2.font.color.rgb = WHITE
                sp2.alignment = PP_ALIGN.CENTER
                sp2.space_before = Pt(20)

        elif data["type"] == "terminal":
            # Title
            tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
            tf = tx_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = CYAN_NEON

            # Command list box
            c_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(11.73), Inches(4.8))
            c_box.fill.solid()
            c_box.fill.fore_color.rgb = CARD_BG
            c_box.line.color.rgb = CYAN_NEON

            ctf = c_box.text_frame
            ctf.word_wrap = True
            ctf.margin_left = Inches(0.4)
            ctf.margin_top = Inches(0.4)

            for idx, item in enumerate(data["commands"]):
                p = ctf.paragraphs[0] if idx == 0 else ctf.add_paragraph()
                p.text = f"> {item['cmd']}"
                p.font.size = Pt(18)
                p.font.bold = True
                p.font.color.rgb = EMERALD
                if idx > 0: p.space_before = Pt(14)

                p2 = ctf.add_paragraph()
                p2.text = f"   {item['desc']}"
                p2.font.size = Pt(14)
                p2.font.color.rgb = LIGHT_GRAY

        elif data["type"] == "roadmap":
            # Title
            tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
            tf = tx_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = CYAN_NEON

            for i, st in enumerate(data["steps"]):
                s_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0 + i * 1.6), Inches(11.73), Inches(1.3))
                s_box.fill.solid()
                s_box.fill.fore_color.rgb = CARD_BG
                s_box.line.color.rgb = CYAN_NEON if i == 0 else EMERALD if i == 1 else YELLOW

                stf = s_box.text_frame
                stf.word_wrap = True
                stf.margin_left = Inches(0.3)
                stf.margin_top = Inches(0.2)

                sp = stf.paragraphs[0]
                sp.text = st["phase"]
                sp.font.size = Pt(18)
                sp.font.bold = True
                sp.font.color.rgb = WHITE

                sp2 = stf.add_paragraph()
                sp2.text = st["detail"]
                sp2.font.size = Pt(14)
                sp2.font.color.rgb = LIGHT_GRAY

        elif data["type"] == "conclusion":
            # Title
            tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
            tf = tx_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = CYAN_NEON

            tx_box2 = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.8))
            tf2 = tx_box2.text_frame
            tf2.word_wrap = True

            for idx, item in enumerate(data["bullets"]):
                p = tf2.paragraphs[0] if idx == 0 else tf2.add_paragraph()
                p.text = item
                p.font.size = Pt(22)
                p.font.bold = True
                p.font.color.rgb = WHITE
                p.space_before = Pt(20)

    output_path = os.path.join(PROJECT_ROOT, "CYVANTIS_Startup_Sprint_2026.pptx")
    prs.save(output_path)
    print(f"PowerPoint file saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
